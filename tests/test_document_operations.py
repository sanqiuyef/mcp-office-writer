from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches

from mcp_office_writer.document_ops import batch_edit_office_document, inspect_office_document
from mcp_office_writer.template_merge import merge_office_template


class DocumentOperationsTest(unittest.TestCase):
    def test_word_node_paths_and_transactional_batch_edit(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "report.docx"
            document = Document()
            document.add_paragraph("First paragraph")
            document.add_paragraph("Move me")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Old cell"
            document.save(path)

            inspected = inspect_office_document(str(path))
            self.assertTrue(inspected["success"], inspected)
            children = inspected["root"]["children"][0]["children"]
            self.assertEqual(children[0]["path"], "/body/p[1]")
            self.assertEqual(children[2]["path"], "/body/tbl[1]")

            result = batch_edit_office_document(str(path), [
                {"op": "set", "path": "/body/p[1]", "props": {"text": "Updated paragraph", "bold": True}},
                {"op": "set", "path": "/body/tbl[1]/row[1]/cell[1]", "props": {"text": "New cell"}},
                {"op": "add", "path": "/body", "element_type": "heading", "props": {"text": "Added heading", "level": 2}},
                {"op": "move", "path": "/body/p[2]", "target_path": "/body/p[1]", "position": "before"},
                {"op": "replace", "path": "/body", "props": {"find": "Added", "replace": "Final"}},
            ])
            self.assertTrue(result["success"], result)
            edited = Document(path)
            self.assertEqual(edited.paragraphs[0].text, "Move me")
            self.assertEqual(edited.paragraphs[1].text, "Updated paragraph")
            self.assertIn("Final heading", [paragraph.text for paragraph in edited.paragraphs])
            self.assertEqual(edited.tables[0].cell(0, 0).text, "New cell")

    def test_failed_batch_does_not_save_partial_word_changes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "atomic.docx"
            document = Document()
            document.add_paragraph("Original")
            document.save(path)
            before = path.read_bytes()

            result = batch_edit_office_document(str(path), [
                {"op": "set", "path": "/body/p[1]", "props": {"text": "Changed in memory"}},
                {"op": "set", "path": "/body/p[99]", "props": {"text": "Invalid"}},
            ])
            self.assertFalse(result["success"])
            self.assertFalse(result["saved"])
            self.assertEqual(result["error"]["code"], "not_found")
            self.assertEqual(path.read_bytes(), before)

    def test_excel_paths_ranges_and_sheet_operations(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "data.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["A1"] = "old"
            worksheet["A1"].font = Font(underline="single")
            workbook.save(path)
            workbook.close()

            result = batch_edit_office_document(str(path), [
                {"op": "set", "path": "/Data/A1:B2", "props": {"values": [[1, 2], [3, 4]], "bold": True}},
                {"op": "add", "path": "/", "element_type": "sheet", "props": {"name": "Summary"}},
                {"op": "set", "path": "/Summary/A1", "props": {"value": "Total"}},
                {"op": "move", "path": "/Summary", "target_path": "/Data", "position": "before"},
            ])
            self.assertTrue(result["success"], result)
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary", "Data"])
                self.assertEqual(workbook["Data"]["B2"].value, 4)
                self.assertTrue(workbook["Data"]["A1"].font.bold)
                self.assertEqual(workbook["Data"]["A1"].font.underline, "single")
                self.assertEqual(workbook["Summary"]["A1"].value, "Total")
            finally:
                workbook.close()

            inspected = inspect_office_document(str(path))
            self.assertTrue(inspected["success"], inspected)
            paths = [child["path"] for child in inspected["root"]["children"]]
            self.assertEqual(paths, ["/Summary", "/Data"])

    def test_powerpoint_paths_add_set_remove_and_move(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "deck.pptx"
            presentation = Presentation()
            first = presentation.slides.add_slide(presentation.slide_layouts[6])
            first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "First"
            second = presentation.slides.add_slide(presentation.slide_layouts[6])
            second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Second"
            presentation.save(path)

            result = batch_edit_office_document(str(path), [
                {"op": "set", "path": "/slide[1]/shape[1]", "props": {"text": "Updated", "font_size": 18}},
                {"op": "add", "path": "/slide[1]", "element_type": "textbox", "props": {"text": "Added", "x": 1, "y": 2}},
                {"op": "remove", "path": "/slide[2]/shape[1]"},
                {"op": "move", "path": "/slide[2]", "target_path": "/slide[1]", "position": "before"},
            ])
            self.assertTrue(result["success"], result)
            presentation = Presentation(path)
            self.assertEqual(len(presentation.slides), 2)
            self.assertEqual(len(presentation.slides[0].shapes), 0)
            self.assertEqual(presentation.slides[1].shapes[0].text, "Updated")
            self.assertEqual(presentation.slides[1].shapes[1].text, "Added")

    def test_docx_template_merge_handles_split_runs_and_unresolved_keys(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            output = Path(td) / "output.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Hello {{user.")
            paragraph.add_run("name}}!")
            document.add_paragraph("Keep {{missing}} visible")
            document.save(template)

            result = merge_office_template(str(template), str(output), {"user": {"name": "Alice"}})
            self.assertTrue(result["success"], result)
            self.assertEqual(result["replaced_count"], 1)
            self.assertEqual(result["used_keys"], ["user.name"])
            self.assertEqual(result["unresolved_placeholders"], ["missing"])
            merged = Document(output)
            self.assertEqual(merged.paragraphs[0].text, "Hello Alice!")
            self.assertIn("{{missing}}", merged.paragraphs[1].text)

    def test_excel_and_powerpoint_template_merge_nested_values(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            excel_template = Path(td) / "template.xlsx"
            excel_output = Path(td) / "output.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "{{items[0].title}}"
            workbook.save(excel_template)
            workbook.close()
            excel_result = merge_office_template(
                str(excel_template), str(excel_output), {"items": [{"title": "第一项"}]}
            )
            self.assertTrue(excel_result["success"], excel_result)
            workbook = load_workbook(excel_output)
            try:
                self.assertEqual(workbook.active["A1"].value, "第一项")
            finally:
                workbook.close()

            ppt_template = Path(td) / "template.pptx"
            ppt_output = Path(td) / "output.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Project: {{project.name}}"
            presentation.save(ppt_template)
            ppt_result = merge_office_template(
                str(ppt_template), str(ppt_output), {"project": {"name": "共享 MCP"}}
            )
            self.assertTrue(ppt_result["success"], ppt_result)
            merged = Presentation(ppt_output)
            self.assertEqual(merged.slides[0].shapes[0].text, "Project: 共享 MCP")

    def test_structured_errors_include_code_path_and_hint(self) -> None:
        result = inspect_office_document("missing.docx")
        self.assertFalse(result["success"])
        self.assertEqual(result["error"]["code"], "invalid_input")
        self.assertEqual(result["error"]["path"], "missing.docx")


if __name__ == "__main__":
    unittest.main()
