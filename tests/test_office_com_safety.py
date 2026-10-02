from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook

from mcp_office_writer.excel_writer import edit_excel_workbook
from mcp_office_writer.models import SectionItem, SheetDef
from mcp_office_writer.office_com import ActiveOfficeDocumentError, _validate_excel_sheets, _write_sheet
from mcp_office_writer.word_writer import edit_word_document


class _FakeWorksheet:
    def __init__(self, name: str = "Sheet1") -> None:
        self.Name = name
        self.cells: dict[tuple[int, int], _FakeCell] = {}

    def Cells(self, row: int, column: int):
        return self.cells.setdefault((row, column), _FakeCell())


class _FakeCell:
    Value = None


class _FakeWorksheets:
    def __init__(self, names: list[str]) -> None:
        self.items = [_FakeWorksheet(name) for name in names]

    @property
    def Count(self) -> int:
        return len(self.items)

    def Item(self, index: int) -> _FakeWorksheet:
        return self.items[index - 1]

    def __call__(self, index: int) -> _FakeWorksheet:
        return self.Item(index)

    def Add(self, After: _FakeWorksheet) -> _FakeWorksheet:
        worksheet = _FakeWorksheet()
        self.items.append(worksheet)
        return worksheet


class _FakeWorkbook:
    def __init__(self, names: list[str]) -> None:
        self.Worksheets = _FakeWorksheets(names)


class OfficeComSafetyTest(unittest.TestCase):
    def test_com_unavailable_falls_back_to_legacy_word_writer(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "fallback.docx"
            Document().save(target)
            with patch("mcp_office_writer.office_com.com_is_available", return_value=False):
                result = edit_word_document(str(target), [SectionItem(type="paragraph", text="legacy fallback")])
            self.assertTrue(result.success, result.message)
            self.assertIn("legacy fallback", "\n".join(paragraph.text for paragraph in Document(target).paragraphs))

    def test_com_unavailable_falls_back_to_legacy_excel_writer(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "fallback.xlsx"
            workbook = Workbook()
            workbook.save(target)
            workbook.close()
            with patch("mcp_office_writer.office_com.com_is_available", return_value=False):
                result = edit_excel_workbook(str(target), [SheetDef(name="Added", data=[["value"]])])
            self.assertTrue(result.success, result.message)
            workbook = load_workbook(target)
            self.assertIn("Added", workbook.sheetnames)
            workbook.close()

    def test_active_excel_formula_is_written_as_literal(self) -> None:
        workbook = _FakeWorkbook(["Existing"])
        _write_sheet(workbook, SheetDef(name="Added", data=[["=1+1"]]))
        self.assertEqual(workbook.Worksheets.Item(2).cells[(1, 1)].Value, "'=1+1")

    def test_active_excel_preflight_rejects_existing_and_duplicate_names(self) -> None:
        workbook = _FakeWorkbook(["Existing"])
        with self.assertRaisesRegex(ActiveOfficeDocumentError, "已存在"):
            _validate_excel_sheets(workbook, [SheetDef(name="existing")])
        with self.assertRaisesRegex(ActiveOfficeDocumentError, "已存在"):
            _validate_excel_sheets(workbook, [SheetDef(name="New"), SheetDef(name="new")])

    def test_active_excel_preflight_rejects_invalid_name_before_mutation(self) -> None:
        workbook = _FakeWorkbook(["Existing"])
        with self.assertRaisesRegex(ActiveOfficeDocumentError, "无效"):
            _validate_excel_sheets(workbook, [SheetDef(name="bad/name")])

