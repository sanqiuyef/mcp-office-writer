from __future__ import annotations

import tempfile
import unittest
from base64 import b64decode
from pathlib import Path

from docx import Document

from mcp_office_writer.checker import validate_word_document
from mcp_office_writer.markdown_writer import create_word_from_markdown, markdown_to_word_content
from mcp_office_writer.word_writer import MAX_RUN_TEXT_CHARS


class MarkdownWordTest(unittest.TestCase):
    def test_parser_handles_report_structures(self) -> None:
        content = markdown_to_word_content(
            "# 标题\n\n正文 **强调**。\n\n- 项目一\n- 项目二\n\n| A | B |\n| --- | --- |\n| 甲 | 乙 |\n"
        )
        self.assertEqual([item.type for item in content.sections], [
            "heading", "paragraph", "bullet_list", "table",
        ])
        self.assertEqual(content.sections[1].text, "正文 强调。")
        self.assertEqual(content.sections[3].table.headers, ["A", "B"])

    def test_long_markdown_is_written_without_json_truncation(self) -> None:
        marker_start = "长文起点-房县"
        marker_end = "长文终点-完整保留"
        long_paragraph = "这是用于验证 Word 长文本导出的正文。" * 1_500
        markdown = f"# {marker_start}\n\n{long_paragraph}\n\n## 结论\n\n{marker_end}\n"

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "report.md"
            output = Path(temp_dir) / "report.docx"
            source.write_text(markdown, encoding="utf-8")
            result = create_word_from_markdown(str(source), str(output), scene="report")
            self.assertTrue(result.success, result.message)

            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn(marker_start, text)
            self.assertIn(marker_end, text)
            self.assertIn(long_paragraph, text)
            self.assertTrue(all(
                len(run.text) <= MAX_RUN_TEXT_CHARS
                for paragraph in document.paragraphs for run in paragraph.runs
            ))

    def test_missing_template_returns_a_clear_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "report.md"
            output = Path(temp_dir) / "report.docx"
            source.write_text("# 标题\n\n正文", encoding="utf-8")
            result = create_word_from_markdown(
                str(source), str(output), template=str(Path(temp_dir) / "missing.dotx")
            )
            self.assertFalse(result.success)
            self.assertIn("文件不存在", result.message)

    def test_markdown_embeds_local_images_and_preserves_missing_image_notice(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image_path = root / "chart.png"
            image_path.write_bytes(b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
                "xHITdwAAAABJRU5ErkJggg=="
            ))
            source = root / "report.md"
            source.write_text("# 报告\n\n![趋势图](chart.png)\n\n![缺图](missing.png)", encoding="utf-8")
            output = root / "report.docx"

            result = create_word_from_markdown(str(source), str(output), scene=None)
            self.assertTrue(result.success, result.message)
            document = Document(output)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertIn("[图片未嵌入: 缺图]", "\n".join(p.text for p in document.paragraphs))

    def test_validator_flags_unresolved_delivery_markers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "draft.docx"
            document = Document()
            document.add_paragraph("TODO: replace the draft conclusion")
            document.save(path)

            report = validate_word_document(str(path))
            marker_check = next(check for check in report.checks if check.check_name == "遗留占位标记")
            self.assertFalse(marker_check.passed)
            self.assertEqual(marker_check.severity, "warn")

    def test_markdown_conversion_preserves_a_word_template(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "report.md"
            template = root / "template.docx"
            output = root / "report.docx"
            source.write_text("# 报告标题\n\n模板导出正文", encoding="utf-8")
            template_document = Document()
            template_document.sections[0].header.paragraphs[0].text = "模板页眉标识"
            template_document.save(template)

            result = create_word_from_markdown(str(source), str(output), template=str(template))
            self.assertTrue(result.success, result.message)
            document = Document(output)
            self.assertEqual(document.sections[0].header.paragraphs[0].text, "模板页眉标识")
            self.assertIn("模板导出正文", "\n".join(p.text for p in document.paragraphs))
