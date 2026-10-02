"""office_convert（dsoffice 封装层）单元测试。

全部 mock 掉 dsoffice 子进程，不要求测试机安装 @deepseek-ai/libreoffice-kit；
真实 CLI 行为由 tmp_dsoffice_test 手工冒烟覆盖。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

from mcp_office_writer import office_convert as oc


def _make_docx(path: Path, text: str = "转换测试") -> Path:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


class FakeRun:
    """替身 _run_dsoffice：记录调用参数，按需落盘输出并返回 JSON payload。

    严格执行 dsoffice 的真实契约：调用时输出文件/目录必须不存在，
    输出目录由 CLI 自己创建——用于拦截封装层预先 mkdir 之类的契约违背。
    """

    def __init__(self, payload: dict | None = None, create_output: bool = True):
        self.payload = payload if payload is not None else {
            "backend": "native", "missingFonts": ["Courier"],
        }
        self.create_output = create_output
        self.calls: list[list[str]] = []

    def __call__(self, cli_path: str, args: list[str], timeout: float) -> dict:
        self.calls.append(args)
        if self.create_output:
            if "--output-dir" in args:
                out = Path(args[args.index("--output-dir") + 1])
                if out.exists():
                    raise RuntimeError(f"EEXIST: file already exists, mkdir '{out}'")
                out.mkdir(parents=True)
            else:
                out = Path(args[args.index("--output") + 1])
                if out.exists():
                    raise RuntimeError(f"EEXIST: file already exists, open '{out}'")
                out.parent.mkdir(parents=True, exist_ok=True)
                out.write_text("%PDF-fake", encoding="utf-8")
        return self.payload


class ConvertToPdfTest(unittest.TestCase):
    def test_convert_with_default_output_and_font_report(self) -> None:
        fake = FakeRun()
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.convert_to_pdf(str(src))
            self.assertTrue(result["success"], msg=result.get("message"))
            expected = (Path(temp_dir) / "报告.pdf").resolve()
            self.assertEqual(Path(result["file_path"]).resolve(), expected)
            self.assertEqual(result["missing_fonts"], ["Courier"])
            self.assertIn("Courier", result["message"])
            # --output 传给 CLI 的应是同一目标
            cli_output = Path(fake.calls[0][fake.calls[0].index("--output") + 1])
            self.assertEqual(cli_output, Path(expected))

    def test_convert_overwrite_removes_existing_output(self) -> None:
        fake = FakeRun()
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            out = Path(temp_dir) / "报告.pdf"
            out.write_text("旧文件内容", encoding="utf-8")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.convert_to_pdf(str(src), str(out), if_exists="overwrite")
            self.assertTrue(result["success"], msg=result.get("message"))
            self.assertEqual(out.read_text(encoding="utf-8"), "%PDF-fake")

    def test_convert_error_mode_raises_when_output_exists(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            out = Path(temp_dir) / "报告.pdf"
            out.write_text("已存在", encoding="utf-8")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", FakeRun()), \
                 self.assertRaises(FileExistsError):
                oc.convert_to_pdf(str(src), str(out), if_exists="error")

    def test_convert_rename_returns_new_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            out = Path(temp_dir) / "报告.pdf"
            out.write_text("已存在", encoding="utf-8")
            fake = FakeRun()
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.convert_to_pdf(str(src), str(out), if_exists="rename")
            self.assertTrue(result["success"], msg=result.get("message"))
            self.assertNotEqual(Path(result["file_path"]), out)
            self.assertTrue(Path(result["file_path"]).exists())

    def test_convert_rejects_unsupported_suffix(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = Path(temp_dir) / "备注.wps"
            src.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError):
                oc.convert_to_pdf(str(src))


class RenderPagesTest(unittest.TestCase):
    def test_render_default_dir_and_payload(self) -> None:
        fake = FakeRun(payload={
            "pageCount": 2,
            "images": [{"index": 1, "page": 1, "path": "p1.png", "width": 816, "height": 1056}],
            "missingFonts": [],
        })
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.render_document_pages(str(src))
            self.assertTrue(result["success"], msg=result.get("message"))
            self.assertEqual(result["page_count"], 2)
            self.assertEqual(result["images"][0]["path"], "p1.png")
            args = fake.calls[0]
            self.assertIn("--pages", args)
            self.assertEqual(args[args.index("--pages") + 1], "all")
            self.assertEqual(args[args.index("--dpi") + 1], "144")
            used_dir = Path(args[args.index("--output-dir") + 1])
            self.assertEqual(used_dir, (Path(temp_dir) / "报告_pages").resolve())

    def test_render_page_list_and_sheet_range_passthrough(self) -> None:
        fake = FakeRun(payload={"pageCount": 0, "images": [], "missingFonts": []})
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "表格.xlsx")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                oc.render_document_pages(str(src), pages="1, 3", dpi=300, sheet="数据", cell_range="A1:D20")
            args = fake.calls[0]
            self.assertEqual(args[args.index("--pages") + 1], "1,3")
            self.assertEqual(args[args.index("--dpi") + 1], "300")
            self.assertEqual(args[args.index("--sheet") + 1], "数据")
            self.assertEqual(args[args.index("--range") + 1], "A1:D20")

    def test_render_rejects_bad_pages(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            with self.assertRaises(ValueError):
                oc.render_document_pages(str(src), pages="1-3")

    def test_render_cell_range_requires_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "表格.xlsx")
            with self.assertRaises(ValueError):
                oc.render_document_pages(str(src), cell_range="A1:D20")

    def test_render_overwrite_removes_existing_dir_before_run(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            stale = Path(temp_dir) / "报告_pages"
            stale.mkdir()
            (stale / "旧图.png").write_text("旧", encoding="utf-8")
            fake = FakeRun()
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.render_document_pages(str(src))
            self.assertTrue(result["success"], msg=result.get("message"))
            self.assertFalse((stale / "旧图.png").exists())

    def test_render_error_mode_keeps_existing_dir(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = _make_docx(Path(temp_dir) / "报告.docx")
            existing = Path(temp_dir) / "报告_pages"
            existing.mkdir()
            (existing / "keep.txt").write_text("保留", encoding="utf-8")
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", FakeRun()), \
                 self.assertRaises(FileExistsError):
                oc.render_document_pages(str(src), if_exists="error")
            self.assertTrue((existing / "keep.txt").exists())


class RecalculateTest(unittest.TestCase):
    def test_recalc_default_target_and_success(self) -> None:
        fake = FakeRun(payload={"backend": "native", "missingFonts": []})
        with tempfile.TemporaryDirectory() as temp_dir:
            src = Path(temp_dir) / "预算.xlsx"
            import openpyxl

            wb = openpyxl.Workbook()
            wb.active["A1"] = 1
            wb.save(str(src))
            with mock.patch.object(oc, "find_dsoffice", return_value="dsoffice-fake"), \
                 mock.patch.object(oc, "_run_dsoffice", fake):
                result = oc.recalculate_excel_formulas(str(src))
            self.assertTrue(result["success"], msg=result.get("message"))
            self.assertEqual(
                Path(result["file_path"]).resolve(),
                (Path(temp_dir) / "预算_recalc.xlsx").resolve(),
            )

    def test_recalc_rejects_non_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            src = Path(temp_dir) / "数据.csv"
            src.write_text("a,b\n1,2", encoding="utf-8")
            with self.assertRaises(ValueError):
                oc.recalculate_excel_formulas(str(src))


class EnvironmentTest(unittest.TestCase):
    def test_find_dsoffice_missing_raises_with_hint(self) -> None:
        with mock.patch.dict(oc.os.environ, {oc.DSOFFICE_ENV_VAR: ""}), \
             mock.patch.object(oc.shutil, "which", return_value=None), \
             self.assertRaises(RuntimeError) as ctx:
            oc.find_dsoffice()
        self.assertIn("npm install", str(ctx.exception))

    def test_find_dsoffice_env_var_precedence(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_cli = Path(temp_dir) / "dsoffice.cmd"
            fake_cli.write_text("echo ok", encoding="utf-8")
            with mock.patch.dict(oc.os.environ, {oc.DSOFFICE_ENV_VAR: str(fake_cli)}):
                self.assertEqual(oc.find_dsoffice(), str(fake_cli))

    def test_run_dsoffice_failure_raises_with_cli_error(self) -> None:
        failed = json.dumps({"code": "failed", "error": "ENOENT: no such file"})
        proc = mock.Mock(returncode=1, stdout=failed, stderr="")
        with mock.patch.object(oc.subprocess, "run", return_value=proc), \
             self.assertRaises(RuntimeError) as ctx:
            oc._run_dsoffice("dsoffice-fake", ["convert"], timeout=10)
        self.assertIn("ENOENT", str(ctx.exception))

    def test_run_dsoffice_timeout_raises(self) -> None:
        import subprocess as sp

        with mock.patch.object(oc.subprocess, "run", side_effect=sp.TimeoutExpired(cmd="dsoffice", timeout=1)), \
             self.assertRaises(RuntimeError) as ctx:
            oc._run_dsoffice("dsoffice-fake", ["convert"], timeout=1)
        self.assertIn("超时", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
