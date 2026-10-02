from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import anyio
from docx import Document
from mcp import ClientSession
from mcp.client.stdio import StdioServerParameters, stdio_client


class McpOfficeEndToEndTest(unittest.TestCase):
    def test_word_excel_and_powerpoint_round_trip_over_mcp(self) -> None:
        async def check() -> None:
            with tempfile.TemporaryDirectory() as temp_dir:
                root = Path(temp_dir)
                params = StdioServerParameters(
                    command=sys.executable,
                    args=["-I", "-X", "utf8", "-m", "mcp_office_writer"],
                    cwd=Path(__file__).resolve().parents[1],
                    encoding="utf-8",
                )
                async with stdio_client(params) as (read_stream, write_stream):
                    async with ClientSession(read_stream, write_stream) as session:
                        await session.initialize()
                        tools = {tool.name for tool in (await session.list_tools()).tools}
                        self.assertIn("create_word_from_markdown", tools)
                        self.assertIn("inspect_office_document", tools)
                        self.assertIn("batch_edit_office_document", tools)
                        self.assertIn("merge_office_template", tools)
                        self.assertIn("get_open_office_documents", tools)
                        self.assertIn("edit_open_word_document", tools)
                        self.assertIn("edit_open_excel_workbook", tools)

                        word = root / "中文报告.docx"
                        word_result = await session.call_tool("create_word_document", {
                            "output_path": str(word),
                            "content": {"sections": [{"type": "paragraph", "text": "中文 Word 内容"}]},
                            "scene": "report",
                        })
                        self.assertFalse(word_result.isError)
                        word_tree = await session.call_tool("inspect_office_document", {"file_path": str(word)})
                        self.assertFalse(word_tree.isError)
                        self.assertIn("/body", word_tree.content[0].text)
                        word_edit = await session.call_tool("batch_edit_office_document", {
                            "file_path": str(word),
                            "operations": [{
                                "op": "replace", "path": "/body",
                                "props": {"find": "中文 Word 内容", "replace": "中文 Word 已更新"},
                            }],
                        })
                        self.assertFalse(word_edit.isError)
                        word_read = await session.call_tool("read_word_document", {"file_path": str(word)})
                        self.assertIn("中文 Word 已更新", word_read.content[0].text)

                        template = root / "模板.docx"
                        merged = root / "合并.docx"
                        template_doc = Document()
                        template_doc.add_paragraph("项目：{{project.name}}")
                        template_doc.save(template)
                        merge_result = await session.call_tool("merge_office_template", {
                            "template_path": str(template),
                            "output_path": str(merged),
                            "data": {"project": {"name": "共享 MCP"}},
                        })
                        self.assertFalse(merge_result.isError)
                        self.assertEqual(Document(merged).paragraphs[0].text, "项目：共享 MCP")

                        excel = root / "数据.xlsx"
                        excel_result = await session.call_tool("create_excel_workbook", {
                            "output_path": str(excel),
                            "sheets": [{"name": "数据", "headers": ["项目"], "data": [["房县"]]}],
                        })
                        self.assertFalse(excel_result.isError)
                        excel_read = await session.call_tool("read_excel_workbook", {"file_path": str(excel)})
                        self.assertIn("房县", excel_read.content[0].text)

                        ppt = root / "汇报.pptx"
                        ppt_result = await session.call_tool("create_powerpoint", {
                            "output_path": str(ppt),
                            "slides": [{"layout": "content", "title": "汇报标题", "body": [{"type": "paragraph", "text": "汇报正文"}]}],
                        })
                        self.assertFalse(ppt_result.isError)
                        ppt_read = await session.call_tool("read_powerpoint", {"file_path": str(ppt)})
                        self.assertIn("汇报标题", ppt_read.content[0].text)
                        self.assertIn("汇报正文", ppt_read.content[0].text)

        anyio.run(check)


if __name__ == "__main__":
    unittest.main()
