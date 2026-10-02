from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
import pymupdf
from pptx import Presentation

from mcp_office_writer import read_excel_workbook as mcp_read_excel_workbook
from mcp_office_writer.checker import validate_word_document
from mcp_office_writer.excel_writer import edit_excel_workbook
from mcp_office_writer.excel_writer import create_excel_workbook
from mcp_office_writer.markdown_writer import _read_markdown
from mcp_office_writer.models import SheetDef
from mcp_office_writer.reader import (
    read_excel_workbook,
    read_pdf_document,
    read_powerpoint,
    read_txt_document,
    read_word_document,
)
from mcp_office_writer.utils import atomic_save, validate_input_file


class SafetyRegressionsTest(unittest.TestCase):
    def test_atomic_save_preserves_existing_file_when_save_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "original.docx"
            target.write_text("original", encoding="utf-8")

            def fail_after_writing(path: str) -> None:
                Path(path).write_text("partial", encoding="utf-8")
                raise RuntimeError("simulated save failure")

            with self.assertRaisesRegex(RuntimeError, "simulated"):
                atomic_save(fail_after_writing, str(target))
            self.assertEqual(target.read_text(encoding="utf-8"), "original")

    def test_excel_formula_is_literal_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "safe.xlsx"
            result = create_excel_workbook(
                str(output), [SheetDef(name="数据", data=[["=1+1"]])]
            )
            self.assertTrue(result.success, result.message)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.active["A1"].value, "'=1+1")
            workbook.close()

    def test_excel_reader_rejects_xls_and_bounds_response(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "large.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "x" * 32000
            workbook.save(workbook_path)
            workbook.close()

            result = read_excel_workbook(str(workbook_path), max_rows_per_sheet=1, max_chars=100)
            self.assertTrue(result["success"])
            self.assertTrue(result["truncated"])
            self.assertLessEqual(len(result["content"]), 100)

            legacy_path = root / "legacy.xls"
            legacy_path.write_bytes(workbook_path.read_bytes())
            legacy_result = read_excel_workbook(str(legacy_path))
            self.assertFalse(legacy_result["success"])
            self.assertIn("不支持", legacy_result["message"])

    def test_pdf_reader_dependency_and_page_count(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "sample.pdf"
            doc = pymupdf.open()
            doc.new_page(width=595, height=842)
            doc.save(str(pdf_path))
            doc.close()

            result = read_pdf_document(str(pdf_path))
            self.assertTrue(result["success"], result["message"])
            self.assertEqual(result["pages_count"], 1)
            self.assertIn("Page 1", result["content"])

    def test_utf16_reading_and_all_reader_response_limits(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            text_path = root / "utf16.txt"
            text_path.write_text("中文 UTF-16 内容", encoding="utf-16")
            text_result = read_txt_document(str(text_path), max_chars=100)
            self.assertTrue(text_result["success"], text_result["message"])
            self.assertEqual(text_result["encoding"], "utf-16")
            self.assertIn("中文", text_result["content"])

            markdown_path = root / "utf16.md"
            markdown_path.write_text("# 中文标题\n\n正文", encoding="utf-16")
            self.assertIn("中文标题", _read_markdown(str(markdown_path)))

            word_path = root / "long.docx"
            document = Document()
            document.add_paragraph("x" * 100)
            document.save(word_path)

            ppt_path = root / "long.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1_000_000, 1_000_000).text = "x" * 100
            presentation.save(ppt_path)

            pdf_path = root / "long.pdf"
            doc = pymupdf.open()
            doc.new_page(width=595, height=842)
            doc.save(str(pdf_path))
            doc.close()

            for reader, path in (
                (read_word_document, word_path),
                (read_powerpoint, ppt_path),
                (read_pdf_document, pdf_path),
                (read_txt_document, text_path),
            ):
                result = reader(str(path), max_chars=5)
                self.assertTrue(result["success"], result["message"])
                self.assertTrue(result["truncated"])
                self.assertLessEqual(len(result["content"]), 5)

    def test_excel_mcp_wrapper_forwards_max_chars(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "large.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "x" * 1000
            workbook.save(workbook_path)
            workbook.close()

            result = mcp_read_excel_workbook(str(workbook_path), max_rows_per_sheet=1, max_chars=40)
            self.assertIsInstance(result, str)
            self.assertIn("内容已截断", result)
            # 正文被截断到 40 字符，元数据行附后，总长保持小体积
            self.assertLess(len(result), 200)

    def test_ooxml_templates_and_validator_use_input_safety_checks(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            for suffix in (".dotx", ".xltx", ".xltm", ".potx"):
                invalid_template = root / f"invalid{suffix}"
                invalid_template.write_text("not an Office archive", encoding="utf-8")
                with self.assertRaises(ValueError):
                    validate_input_file(str(invalid_template), (suffix,))

            invalid_word = root / "invalid.docx"
            invalid_word.write_text("not an Office archive", encoding="utf-8")
            report = validate_word_document(str(invalid_word))
            self.assertEqual(report.errors, 1)
            self.assertEqual(report.checks[0].check_name, "输入文件安全检查")

    def test_xlsm_edits_keep_vba_archive(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "book.xlsm"
            workbook = Workbook()
            workbook.save(workbook_path)
            workbook.close()

            with patch("mcp_office_writer.excel_writer.load_workbook", wraps=load_workbook) as mocked_load:
                result = edit_excel_workbook(str(workbook_path), [])
            self.assertTrue(result.success, result.message)
            self.assertTrue(mocked_load.call_args.kwargs["keep_vba"])

    def test_xlsm_creation_requires_a_macro_template(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            result = create_excel_workbook(
                str(Path(temp_dir) / "macro.xlsm"),
                [SheetDef(name="Sheet1", data=[["value"]])],
            )
            self.assertFalse(result.success)
            self.assertIn(".xlsm", result.message)
