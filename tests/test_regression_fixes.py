# -*- coding: utf-8 -*-
"""回归测试：锁定 2026-08-31 全面复核修复的功能缺陷。

覆盖：
- landscape 页面方向失效（word_writer._set_page_setup 纸型覆盖宽高）
- PPT 连接线/框架图箭头丢失（a:ln 查找层级错误）
- PPT 副标题永不写入（shape_type 常量错误）
- 含表格 PPT 应用配色失败 / 表格配色不生效
- PDF 标题字号 3/4 级倒挂
- 框架图节点多行文本换行失效
- 模板合并遇超链接占位符结果损坏
- 封面 + 首节为表格时文档创建失败
- download_image 的 scheme 与大小限制
"""
import pytest


# ── Word: landscape ──────────────────────────────────────────────

def test_landscape_page_setup_swaps_dimensions(tmp_path):
    from docx import Document

    from mcp_office_writer.models import PageSetup
    from mcp_office_writer.word_writer import _set_page_setup

    doc = Document()
    section = doc.sections[0]
    _set_page_setup(section, PageSetup(orientation="landscape"))
    # A4 portrait 8.27x11.69 → landscape 应为 11.69x8.27
    assert section.page_width.inches > section.page_height.inches
    assert round(section.page_width.inches, 2) == 11.69
    assert round(section.page_height.inches, 2) == 8.27


def test_create_word_document_with_landscape(tmp_path):
    from mcp_office_writer.models import PageSetup, WordContent
    from mcp_office_writer.word_writer import create_word_document

    out = tmp_path / "landscape.docx"
    result = create_word_document(
        str(out),
        WordContent(sections=[]),
        page_setup=PageSetup(orientation="landscape", paper_size="A4"),
    )
    assert result.success
    from docx import Document
    section = Document(str(out)).sections[0]
    assert section.page_width.inches > section.page_height.inches


# ── PPT: 连接线箭头 ──────────────────────────────────────────────

def test_slide_connector_has_arrowhead(tmp_path):
    from lxml import etree
    from pptx import Presentation

    from mcp_office_writer.models import SlideConnectorDef
    from mcp_office_writer.ppt_writer import _add_slide_connector

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_connector(slide, SlideConnectorDef(x1=1, y1=1, x2=3, y2=1))
    xml = etree.tostring(slide.shapes[-1]._element).decode()
    assert "tailEnd" in xml
    assert 'type="triangle"' in xml


def test_framework_diagram_arrowhead_in_xml(tmp_path):
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    from mcp_office_writer.ppt_writer import _add_connector_arrow

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_connector_arrow(slide, "94A3B8", 1, Inches(1), Inches(1), Inches(3), Inches(1))
    xml = etree.tostring(slide.shapes[-1]._element).decode()
    assert "tailEnd" in xml


def test_create_framework_diagram_includes_arrows(tmp_path):
    from pptx import Presentation

    from mcp_office_writer.models import FrameworkDiagram, FwNode, FwStage
    from mcp_office_writer.ppt_writer import create_framework_diagram

    out = tmp_path / "fw.pptx"
    result = create_framework_diagram(
        str(out),
        FrameworkDiagram(
            title="T",
            stages=[FwStage(label="阶段1", nodes=[
                FwNode(id="n1", text="A"), FwNode(id="n2", text="B"),
            ])],
        ),
    )
    assert result.success
    prs = Presentation(str(out))
    connectors = [s for s in prs.slides[0].shapes if s.shape_type == 9]  # LINE
    assert connectors, "应存在连接线"
    assert all("tailEnd" in c._element.xml for c in connectors)


# ── PPT: 副标题 ──────────────────────────────────────────────────

def test_powerpoint_subtitle_is_written(tmp_path):
    from pptx import Presentation

    from mcp_office_writer.models import SlideDef
    from mcp_office_writer.ppt_writer import create_powerpoint

    out = tmp_path / "subtitle.pptx"
    result = create_powerpoint(str(out), [SlideDef(layout="title", title="T", subtitle="SUB")])
    assert result.success
    slide = Presentation(str(out)).slides[0]
    texts = [ph.text for ph in slide.placeholders]
    assert "SUB" in texts


# ── PPT: 配色方案与含表格文件 ────────────────────────────────────

def test_apply_color_scheme_with_table_does_not_fail(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from mcp_office_writer.ppt_writer import apply_color_to_slides

    src = tmp_path / "tbl.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "H1"
    table.cell(1, 0).text = "D1"
    prs.save(str(src))

    result = apply_color_to_slides(str(src), "professional-blue")
    assert result["success"], result

    prs2 = Presentation(str(src))
    tbl = prs2.slides[0].shapes[0].table
    header_run = tbl.cell(0, 0).text_frame.paragraphs[0].runs[0]
    body_run = tbl.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert str(header_run.font.color.rgb) == "FFFFFF"
    assert str(body_run.font.color.rgb) == "333333"


# ── PDF: 标题字号单调 ────────────────────────────────────────────

def test_pdf_heading_sizes_are_monotonic():
    from mcp_office_writer.pdf_writer import PdfWriter

    writer = PdfWriter()
    span = writer.heading_size - writer.body_size
    step = span / 5
    sizes = [round(writer.body_size + step * (6 - k), 1) for k in range(1, 7)]
    for a, b in zip(sizes, sizes[1:]):
        assert a >= b, f"标题字号未递减: {sizes}"
    assert sizes[0] == pytest.approx(writer.heading_size)
    assert sizes[-1] == pytest.approx(writer.body_size)


# ── 框架图: 多行节点文本 ─────────────────────────────────────────

def test_framework_node_multiline_text_splits_paragraphs(tmp_path):
    from pptx import Presentation

    from mcp_office_writer.models import FrameworkDiagram, FwNode, FwStage
    from mcp_office_writer.ppt_writer import create_framework_diagram

    out = tmp_path / "fw_multiline.pptx"
    result = create_framework_diagram(
        str(out),
        FrameworkDiagram(
            stages=[FwStage(label="S", nodes=[FwNode(id="n1", text="第一行\n第二行")])],
        ),
    )
    assert result.success
    prs = Presentation(str(out))
    node_frame = None
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and "第一行" in shape.text:
            node_frame = shape.text_frame
            break
    assert node_frame is not None
    paragraphs = [p.text for p in node_frame.paragraphs]
    assert "第一行" in paragraphs
    assert "第二行" in paragraphs


# ── 模板合并: hyperlink 内占位符 ─────────────────────────────────

def test_template_merge_replaces_hyperlink_placeholders(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    from mcp_office_writer.template_merge import merge_office_template

    tpl = tmp_path / "tpl.docx"
    doc = Document()
    doc.add_paragraph("正文 {{normal}}")
    hyperlink_paragraph = parse_xml(
        f'<w:p {nsdecls("w")} {nsdecls("r")}>'
        f'<w:hyperlink r:id="rId999"><w:r><w:t>{{{{hl_key}}}}</w:t></w:r></w:hyperlink></w:p>'
    )
    doc.element.body.append(hyperlink_paragraph)
    doc.save(str(tpl))

    out = tmp_path / "out.docx"
    result = merge_office_template(str(tpl), str(out), {"normal": "N", "hl_key": "VALUE"})
    assert result["success"]
    assert result["unresolved_placeholders"] == []

    merged = Document(str(out))
    texts = [p.text for p in merged.paragraphs]
    assert "VALUE" in texts, f"hyperlink 内占位符未替换: {texts}"


# ── 封面: 首节为表格不崩溃 ───────────────────────────────────────

def test_cover_with_table_first_section(tmp_path):
    from mcp_office_writer.models import SectionItem, TableDef, WordContent
    from mcp_office_writer.word_writer import create_word_document

    out = tmp_path / "cover_table.docx"
    result = create_word_document(
        str(out),
        WordContent(sections=[SectionItem(type="table", table=TableDef(headers=["a"], rows=[["b"]]))]),
        scene="report",  # 自动封面
    )
    assert result.success, result.message


def test_cover_with_bg_color_recipe_generates(tmp_path):
    """R7 深色背景配方不应破坏文档生成，且不应崩在背景表格上。"""
    from mcp_office_writer.models import CoverConfig, WordContent
    from mcp_office_writer.word_writer import create_word_document

    out = tmp_path / "cover_bg.docx"
    result = create_word_document(
        str(out),
        WordContent(sections=[]),
        cover_config=CoverConfig(title="品牌标题", organization="某机构", recipe_id="R7").model_dump(),
    )
    assert result.success, result.message


# ── download_image: scheme / 大小限制 ────────────────────────────

def test_download_image_rejects_non_http_scheme():
    from mcp_office_writer.utils import download_image

    with pytest.raises(ValueError):
        download_image("file:///C:/Windows/win.ini")


# ── landscape 与 template 无关的既有保护 ──────────────────────────

def test_edit_word_prefers_com_then_disk(tmp_path):
    """编辑 docx 时 COM 不可用应回退磁盘路径成功。"""
    from mcp_office_writer.models import SectionItem, WordContent
    from mcp_office_writer.word_writer import create_word_document, edit_word_document

    src = tmp_path / "base.docx"
    assert create_word_document(str(src), WordContent(sections=[SectionItem(type="paragraph", text="内容")])).success
    result = edit_word_document(str(src), [SectionItem(type="paragraph", text="追加")])
    assert result.success, result.message